import os
import openpyxl
from docx import Document

SAMPLE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_files")
os.makedirs(SAMPLE_DIR, exist_ok=True)

def generate_sample_excel():
    """Generates a realistic multi-unit question bank Excel workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Question Bank"

    headers = ["Sl.No", "Unit", "Question Statement", "Marks", "Question Type", "Difficulty", "Blooms Level", "Course Outcome", "Key Points / Rubric"]
    ws.append(headers)

    sample_questions = [
        (1, 1, "Define a Deterministic Finite Automaton (DFA) with its formal 5-tuple specification.", 2, "Short Answer", "Easy", "Remember", "CO1", "5-tuple (Q, Sigma, delta, q0, F): 2 Marks"),
        (2, 1, "Construct a DFA that accepts all strings over {0,1} ending with the substring '101'.", 8, "Long Answer", "Medium", "Apply", "CO1", "State transition table: 4 Marks, Transition diagram: 4 Marks"),
        (3, 1, "Convert the given Non-Deterministic Finite Automaton (NFA) with epsilon transitions to an equivalent DFA using subset construction.", 13, "Long Answer", "Hard", "Apply", "CO1", "Epsilon-closure calculation: 4 Marks, Subset state mapping: 5 Marks, Minimized DFA: 4 Marks"),
        (4, 2, "State the pumping lemma for regular languages and explain its significance.", 2, "Short Answer", "Easy", "Remember", "CO2", "Statement of pumping lemma: 1 Mark, Proof application: 1 Mark"),
        (5, 2, "Prove that the language L = {0^n 1^n | n >= 1} is not regular using the Pumping Lemma.", 8, "Long Answer", "Medium", "Analyze", "CO2", "Assumption & string selection: 3 Marks, Decomposition xyz: 3 Marks, Contradiction proof: 2 Marks"),
        (6, 2, "Design a Context-Free Grammar (CFG) for the language of balanced parentheses and derive the string '((()))'.", 13, "Long Answer", "Hard", "Create", "CO2", "Production rules: 5 Marks, Leftmost and rightmost derivations: 4 Marks, Parse tree: 4 Marks"),
        (7, 3, "Differentiate between Deterministic and Non-Deterministic Pushdown Automata (PDA).", 2, "Short Answer", "Easy", "Understand", "CO3", "Acceptance by empty store vs final state: 1 Mark, Expressiveness difference: 1 Mark"),
        (8, 3, "Construct a Pushdown Automaton (PDA) accepting the language L = {w w^R | w in {a, b}*}.", 13, "Long Answer", "Hard", "Apply", "CO3", "Stack operations: 4 Marks, Transition functions: 5 Marks, Step trace: 4 Marks"),
        (9, 4, "Define the formal specification of a standard single-tape Turing Machine.", 2, "Short Answer", "Easy", "Remember", "CO4", "7-tuple definition: 2 Marks"),
        (10, 4, "Design a Turing Machine that computes the 1's complement of a binary string and write its transition table.", 13, "Long Answer", "Hard", "Create", "CO4", "State logic: 4 Marks, Transition table: 5 Marks, Tape movement trace: 4 Marks"),
        (11, 5, "Define the Post Correspondence Problem (PCP) and state whether it is decidable.", 2, "Short Answer", "Easy", "Remember", "CO5", "PCP definition: 1.5 Marks, Undecidability: 0.5 Mark"),
        (12, 5, "Explain the Halting Problem of Turing Machines and prove that it is undecidable using diagonalization.", 15, "Essay Questions", "Hard", "Evaluate", "CO5", "Problem statement: 3 Marks, Reduction & Diagonalization proof: 8 Marks, Implications: 4 Marks")
    ]

    for row in sample_questions:
        ws.append(list(row))

    excel_path = os.path.join(SAMPLE_DIR, "CS3452_Theory_Of_Computation_Bank.xlsx")
    wb.save(excel_path)
    print(f"Sample Excel saved to: {excel_path}")

def generate_sample_docx_paper():
    """Generates a sample university examination paper pattern in Word (.docx) format."""
    doc = Document()
    doc.add_heading("ANNA UNIVERSITY AUTONOMOUS EXAMINATIONS - MODEL PAPER", level=1)
    doc.add_paragraph("Department of Computer Science and Engineering")
    doc.add_paragraph("Subject: CS3452 - Theory of Computation | Time: 3 Hours | Max Marks: 100")
    doc.add_paragraph("----------------------------------------------------------------------")

    doc.add_heading("PART - A (10 x 2 = 20 Marks)", level=2)
    doc.add_paragraph("Answer ALL Questions")
    for i in range(1, 11):
        doc.add_paragraph(f"{i}. Sample short question on theoretical computer science. (CO{((i-1)//2)+1}) [RE] (2 Marks)")

    doc.add_heading("PART - B (5 x 13 = 65 Marks)", level=2)
    doc.add_paragraph("Answer ALL Questions (Either / Or Choice)")
    for i in range(11, 16):
        doc.add_paragraph(f"{i}. (a) Explain the construction of finite automata and regular expressions. (CO{i-10}) [AP] (13 Marks)")
        doc.add_paragraph("                           (OR)")
        doc.add_paragraph(f"    (b) Design an equivalent state minimization algorithm and trace execution. (CO{i-10}) [AN] (13 Marks)")

    doc.add_heading("PART - C (1 x 15 = 15 Marks)", level=2)
    doc.add_paragraph("Application / Comprehensive Evaluation (Either / Or Choice)")
    doc.add_paragraph("16. (a) Critically analyze the computational complexity boundaries between P, NP, and NP-Complete problems with real-world reductions. (CO5) [EV] (15 Marks)")
    doc.add_paragraph("                           (OR)")
    doc.add_paragraph("    (b) Formulate a cellular automaton model and evaluate its equivalence to universal computation. (CO5) [CR] (15 Marks)")

    docx_path = os.path.join(SAMPLE_DIR, "Sample_University_Exam_Pattern.docx")
    doc.save(docx_path)
    print(f"Sample DOCX saved to: {docx_path}")

def generate_sample_syllabus():
    """Generates a sample syllabus text document."""
    txt_path = os.path.join(SAMPLE_DIR, "AI3001_Syllabus.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("""SUBJECT CODE: AI3001
SUBJECT NAME: ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING
REGULATION: 2023 AUTONOMOUS

UNIT I: PROBLEM SOLVING & SEARCH ALGORITHMS
Introduction to AI, Agents and Environments, PEAS specification. Uninformed search strategies: BFS, DFS, Uniform Cost Search. Informed search: A* search, Heuristics admissibility and consistency, Hill Climbing, Adversarial search: Minimax, Alpha-Beta pruning.

UNIT II: KNOWLEDGE REPRESENTATION & REASONING
Logical Agents, Propositional Logic, First-Order Logic (FOL), Syntax and Semantics, Forward and Backward Chaining, Resolution refutation, Unification algorithm, Ontological engineering, Handling uncertainty with Bayesian Networks.

UNIT III: SUPERVISED & UNSUPERVISED LEARNING
Machine Learning taxonomy. Linear regression, Logistic regression. Decision Trees, ID3 algorithm, Entropy and Information Gain. Naive Bayes classifier. Support Vector Machines. Clustering: K-Means, Hierarchical clustering, Principal Component Analysis (PCA).

UNIT IV: DEEP LEARNING & NEURAL NETWORKS
Biological vs Artificial Neurons, Perceptron, Activation functions. Multi-Layer Perceptron (MLP), Backpropagation algorithm with gradient descent. Regularization techniques: Dropout, L1/L2 penalties. Convolutional Neural Networks (CNN) for image analysis.

UNIT V: REINFORCEMENT LEARNING & ETHICS
Markov Decision Process (MDP), Bellman Optimality Equations, Temporal Difference Learning, Q-Learning. Ethical challenges in AI: Bias, Fairness, Interpretability, and Responsible AI deployments.
""")
    print(f"Sample Syllabus saved to: {txt_path}")

if __name__ == "__main__":
    generate_sample_excel()
    generate_sample_docx_paper()
    generate_sample_syllabus()
